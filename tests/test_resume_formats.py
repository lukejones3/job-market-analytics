"""Exercise real parsers so a clean install supports every advertised format."""
import io

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from python.resume.parser import parse_resume


def test_pdf_resume_extracts_real_page_text():
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
        NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):
        DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    content = DecodedStreamObject()
    content.set_data(b'BT /F1 12 Tf 72 720 Td (Python SQL Data Analyst) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(content)
    stream = io.BytesIO()
    writer.write(stream)
    assert 'Python SQL Data Analyst' in parse_resume(stream.getvalue(), 'synthetic.pdf')


def test_docx_resume_includes_tables_and_headers():
    document = Document()
    document.add_paragraph('Data analyst with Python')
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'SQL PostgreSQL'
    document.sections[0].header.paragraphs[0].text = 'Analytics portfolio'
    stream = io.BytesIO()
    document.save(stream)
    text = parse_resume(stream.getvalue(), 'synthetic.docx')
    assert all(part in text for part in ['Data analyst', 'SQL PostgreSQL', 'Analytics portfolio'])


def test_plain_text_resume_normalizes_spacing_without_losing_lines():
    assert parse_resume('Python\u00a0  SQL\nAnalytics'.encode(), 'synthetic.txt') == 'Python SQL\nAnalytics'
